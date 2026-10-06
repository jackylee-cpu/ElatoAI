#!/usr/bin/env python3
"""Production-readiness tracker.

docs/production-readiness.xlsx is the source of truth. Every read and write goes
through this script so the update log, the dashboard sheet and the rule that
"done" means "tested and passed" stay consistent with each other.

Priorities, phases and the starting task list live in assets/seed.json, so the
same script serves another project by editing that one file.

    tracker.py status                       progress report (markdown)
    tracker.py show P0-02                   one task or question in full, with history
    tracker.py list [--priority P0] [--phase 1] [--status 進行中]
    tracker.py start P0-02
    tracker.py update P0-02 --status 已完成 --result 通過 \
        --evidence "anon GET /rest/v1/conversations -> 0 rows" --commit abc1234
    tracker.py update P0-02 --result 失敗 --evidence "..."
    tracker.py update P1-08 --deps "Q1, R-01" --phase 3
    tracker.py answer Q2 --text "..."
    tracker.py ask --question "..." --why "..." --affects "R-02"
    tracker.py add --priority R --phase 2 --title "..." [--category --risk --where --accept --test --deps]
    tracker.py sync                         bring seed.json additions and the latest
                                            layout into an existing workbook
    tracker.py init [--force]               build a fresh workbook from seed.json
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import math
import os
import re
import shutil
import sys
import tempfile
import unicodedata
import zipfile
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import CellIsRule, DataBarRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover - environment problem, not a bug
    sys.exit("需要 openpyxl：python3 -m pip install openpyxl")

SKILL_DIR = Path(__file__).resolve().parent.parent
REPO_ROOT = SKILL_DIR.parents[2]  # <repo>/.claude/skills/prod-readiness
DEFAULT_FILE = REPO_ROOT / "docs" / "production-readiness.xlsx"
SEED_FILE = SKILL_DIR / "assets" / "seed.json"

DASH, TASKS, QUESTIONS, LOG = "總覽", "任務清單", "待確認問題", "更新記錄"

# (key, header, column width). Headers are how columns are found on read, so a
# user reordering or inserting columns in Excel does not break anything.
TASK_COLS = [
    ("id", "編號", 8),
    ("priority", "優先級", 7),
    ("phase", "階段", 12),
    ("category", "類別", 13),
    ("title", "任務", 30),
    ("status", "狀態", 9),
    ("result", "測試結果", 10),
    ("done_date", "完成日期", 12),
    ("risk", "問題 / 風險", 46),
    ("where", "相關位置", 42),
    ("accept", "驗收標準", 46),
    ("test", "測試方法", 46),
    ("deps", "依賴", 12),
    ("evidence", "測試證據", 44),
    ("commit", "Commit", 10),
    ("note", "備註", 32),
]
QUESTION_COLS = [
    ("id", "編號", 6),
    ("question", "問題", 36),
    ("why", "為什麼重要", 46),
    ("affects", "影響的任務", 22),
    ("status", "狀態", 9),
    ("answer", "回答", 46),
    ("answered_on", "回答日期", 12),
]
LOG_COLS = [
    ("time", "時間", 17),
    ("id", "編號", 8),
    ("title", "任務 / 問題", 30),
    ("old_status", "舊狀態", 9),
    ("new_status", "新狀態", 9),
    ("result", "測試結果", 10),
    ("evidence", "測試證據", 50),
    ("commit", "Commit", 10),
    ("note", "備註", 36),
]

STATUSES = ["未開始", "進行中", "已完成", "阻塞", "不適用"]
RESULTS = ["未測試", "通過", "失敗", "部分通過"]
Q_STATUSES = ["待回覆", "已回覆"]
STATUS_ALIASES = {
    "todo": "未開始", "doing": "進行中", "wip": "進行中", "done": "已完成",
    "blocked": "阻塞", "na": "不適用", "n/a": "不適用",
}
RESULT_ALIASES = {
    "untested": "未測試", "pass": "通過", "passed": "通過",
    "fail": "失敗", "failed": "失敗", "partial": "部分通過",
}


def load_seed() -> dict:
    try:
        return json.loads(SEED_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        sys.exit(f"讀不到 {SEED_FILE}：{exc}")


SEED = load_seed()
PROJECT = SEED.get("project", "")
# key -> (label, when it has to be done by). Order is sort order.
PRIORITIES = {p["key"]: (p["label"], p.get("when", "")) for p in SEED.get("priorities", [])} or {
    "P0": ("安全", "真實用戶使用前必須完成"),
    "P1": ("兒童安全與合規", "上市前必須完成"),
    "P2": ("工程與運維", "規模化前完成"),
}
# name -> (dates, goal). Order is the roadmap order; empty means no phases.
PHASES = {p["name"]: (p.get("dates", ""), p.get("goal", "")) for p in SEED.get("phases", [])}
PRIORITY_COLORS = ["C00000", "C65911", "2F5597", "7030A0", "548235", "806000"]

FONT = "Arial"
BASE_FONT = Font(name=FONT, size=10)
BOLD_FONT = Font(name=FONT, size=10, bold=True)
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", start_color="1F4E79", end_color="1F4E79")
SECTION_FONT = Font(name=FONT, size=12, bold=True, color="1F4E79")
MUTED_FONT = Font(name=FONT, size=9, color="7F7F7F")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER_TOP = Alignment(horizontal="center", vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center")
CENTERED = {"id", "priority", "phase", "status", "result", "done_date", "commit", "deps",
            "answered_on", "old_status", "new_status", "time"}
MAX_ROW = 1000  # validation and formatting cover this many rows


def die(msg: str) -> None:
    sys.exit(f"錯誤：{msg}")


def now_str() -> str:
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M")


def today() -> str:
    return dt.date.today().isoformat()


# ── layout helpers ───────────────────────────────────────────────────────────

def disp_width(text: str) -> int:
    """Columns a string occupies: CJK characters are two wide."""
    return sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in text)


def row_height(pairs) -> float:
    """Estimate a wrapped row's height; Excel does not re-fit rows on open."""
    lines = 1
    for text, width in pairs:
        if text in (None, ""):
            continue
        usable = max(width - 1.5, 1)
        n = sum(max(1, math.ceil(disp_width(p) / usable)) for p in str(text).split("\n"))
        lines = max(lines, n)
    return min(14 * lines + 4, 409)


def col_letter(cols, key: str) -> str:
    return get_column_letter([k for k, _, _ in cols].index(key) + 1)


def write_header(ws, cols) -> None:
    for i, (_, label, width) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=label)
        c.font, c.fill, c.border = HEADER_FONT, HEADER_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 22


def write_cells(ws, row: int, cmap: dict, cols, record: dict) -> None:
    """Write the keys present in `record` and re-fit the row's height."""
    widths = {k: w for k, _, w in cols}
    for key, value in record.items():
        if key not in cmap:
            continue
        c = ws.cell(row=row, column=cmap[key], value=(value if value != "" else None))
        c.font, c.border = BASE_FONT, BORDER
        c.alignment = CENTER_TOP if key in CENTERED else WRAP_TOP
    ws.row_dimensions[row].height = row_height(
        (ws.cell(row=row, column=cmap[k]).value, widths[k]) for k in cmap
    )


def list_validation(ws, col: str, options) -> None:
    dv = DataValidation(
        type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
        showErrorMessage=True, errorTitle="無效的值", error="請從下拉清單選擇",
    )
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{MAX_ROW}")


def value_colors(ws, col: str, spec) -> None:
    rng = f"{col}2:{col}{MAX_ROW}"
    for value, fill, color in spec:
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="equal", formula=[f'"{value}"'],
            fill=PatternFill("solid", start_color=fill, end_color=fill) if fill else None,
            font=Font(name=FONT, color=color, bold=True),
        ))


def style_sheets(wb) -> None:
    """Validation lists, colours and frozen panes. Applied when a workbook is built."""
    t = wb[TASKS]
    list_validation(t, col_letter(TASK_COLS, "priority"), list(PRIORITIES))
    if PHASES:
        list_validation(t, col_letter(TASK_COLS, "phase"), list(PHASES))
    list_validation(t, col_letter(TASK_COLS, "status"), STATUSES)
    list_validation(t, col_letter(TASK_COLS, "result"), RESULTS)
    value_colors(t, col_letter(TASK_COLS, "status"), [
        ("已完成", "C6EFCE", "006100"), ("進行中", "FFEB9C", "9C5700"),
        ("阻塞", "FFC7CE", "9C0006"), ("不適用", "EDEDED", "7F7F7F"),
    ])
    value_colors(t, col_letter(TASK_COLS, "result"), [
        ("通過", "C6EFCE", "006100"), ("失敗", "FFC7CE", "9C0006"),
        ("部分通過", "FFEB9C", "9C5700"),
    ])
    value_colors(t, col_letter(TASK_COLS, "priority"), [
        (key, None, PRIORITY_COLORS[i % len(PRIORITY_COLORS)]) for i, key in enumerate(PRIORITIES)
    ])
    # Keep everything up to and including the task title in view.
    t.freeze_panes = f"{get_column_letter([k for k, _, _ in TASK_COLS].index('title') + 2)}2"

    q = wb[QUESTIONS]
    list_validation(q, col_letter(QUESTION_COLS, "status"), Q_STATUSES)
    value_colors(q, col_letter(QUESTION_COLS, "status"), [
        ("已回覆", "C6EFCE", "006100"), ("待回覆", "FFEB9C", "9C5700"),
    ])
    q.freeze_panes = "C2"
    wb[LOG].freeze_panes = "A2"


def build_workbook(tasks, questions, log) -> Workbook:
    """A fresh workbook in the current layout, holding the given rows."""
    wb = Workbook()
    wb.active.title = TASKS
    for name in (QUESTIONS, LOG):
        wb.create_sheet(name)
    for name, cols, rows, defaults in (
        (TASKS, TASK_COLS, tasks, {"status": "未開始", "result": "未測試"}),
        (QUESTIONS, QUESTION_COLS, questions, {"status": "待回覆"}),
        (LOG, LOG_COLS, log, {}),
    ):
        ws = wb[name]
        write_header(ws, cols)
        cmap = header_map(ws, cols)
        for r, rec in enumerate(rows, start=2):
            full = {k: rec.get(k, "") for k, _, _ in cols}
            for k, v in defaults.items():
                full[k] = full[k] or v
            write_cells(ws, r, cmap, cols, full)
    style_sheets(wb)
    return wb


# ── reading ──────────────────────────────────────────────────────────────────

def header_map(ws, cols, allow_missing: bool = False) -> dict:
    labels = {label: key for key, label, _ in cols}
    cmap = {labels[c.value]: c.column for c in ws[1] if c.value in labels}
    missing = [label for key, label, _ in cols if key not in cmap]
    if missing and not allow_missing:
        die(f"工作表「{ws.title}」缺少欄位：{'、'.join(missing)}。"
            "如果是舊版的表，先執行：tracker.py sync")
    return cmap


def cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M") if (v.hour or v.minute) else v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return str(v).strip()


def read_rows(ws, cols, allow_missing: bool = False):
    cmap = header_map(ws, cols, allow_missing)
    rows = []
    for r in range(2, ws.max_row + 1):
        rec = {k: cell_text(ws.cell(row=r, column=c).value) for k, c in cmap.items()}
        if not any(rec.values()):
            continue
        for key, _, _ in cols:
            rec.setdefault(key, "")
        rec["_row"] = r
        rows.append(rec)
    return rows, cmap


def read_all(wb, allow_missing: bool = False):
    tasks, _ = read_rows(wb[TASKS], TASK_COLS, allow_missing)
    for t in tasks:
        t["id"] = t["id"].upper()
        t["status"] = t["status"] or "未開始"
        t["result"] = t["result"] or "未測試"
    questions, _ = read_rows(wb[QUESTIONS], QUESTION_COLS, allow_missing)
    for q in questions:
        q["id"] = q["id"].upper()
        q["status"] = q["status"] or "待回覆"
    log, _ = read_rows(wb[LOG], LOG_COLS, allow_missing)
    return tasks, questions, log


def id_num(task_id: str) -> int:
    m = re.search(r"(\d+)$", task_id)
    return int(m.group(1)) if m else 0


def order_of(value: str, ordered) -> int:
    keys = list(ordered)
    return keys.index(value) if value in keys else len(keys)


def sort_key(t) -> tuple:
    """Roadmap order: phase first, then priority, then number."""
    return (order_of(t["phase"], PHASES), order_of(t["priority"], PRIORITIES), id_num(t["id"]))


def split_deps(deps: str) -> list[str]:
    return [d.upper() for d in re.split(r"[,\s，、;；]+", deps or "") if d]


def unmet_deps(task, by_id: dict, q_by_id: dict) -> list[str]:
    unmet = []
    for d in split_deps(task["deps"]):
        if d in by_id:
            if by_id[d]["status"] not in ("已完成", "不適用"):
                unmet.append(d)
        elif d in q_by_id:
            if q_by_id[d]["status"] != "已回覆":
                unmet.append(d)
        else:
            unmet.append(f"{d}（找不到）")
    return unmet


def is_open(t) -> bool:
    return t["status"] not in ("已完成", "不適用")


def plan(tasks, questions, log) -> dict:
    """Everything the report and the dashboard show, computed once."""
    by_id = {t["id"]: t for t in tasks}
    q_by_id = {q["id"]: q for q in questions}
    ordered = sorted(tasks, key=sort_key)

    # Most recently completed first: the log knows the order, done_date only
    # knows the day.
    last_done_at = {}
    for i, entry in enumerate(log):
        if entry["new_status"] == "已完成":
            last_done_at[entry["id"].upper()] = i
    done = sorted(
        (t for t in tasks if t["status"] == "已完成"),
        key=lambda t: (last_done_at.get(t["id"], -1), t["done_date"]), reverse=True,
    )

    in_progress = [t for t in ordered if t["status"] == "進行中"]
    ready, waiting = [], []
    for t in ordered:
        if t["status"] != "未開始":
            continue
        unmet = unmet_deps(t, by_id, q_by_id)
        if unmet:
            waiting.append((t, unmet))
        else:
            ready.append(t)

    active = [t for t in tasks if t["status"] != "不適用"]

    def progress(field, key):
        group = [t for t in active if t[field] == key]
        return sum(t["status"] == "已完成" for t in group), len(group)

    return {
        "done": done,
        "in_progress": in_progress,
        "ready": ready,
        "waiting": waiting,
        "blocked": [t for t in ordered if t["status"] == "阻塞"],
        "open_questions": [q for q in questions if q["status"] != "已回覆"],
        "total": (sum(t["status"] == "已完成" for t in active), len(active)),
        "per_priority": {p: progress("priority", p) for p in PRIORITIES},
        "per_phase": {ph: progress("phase", ph) for ph in PHASES},
        "current_phase": next(
            (ph for ph in PHASES if any(t["phase"] == ph and is_open(t) for t in tasks)), None),
        "last_update": log[-1]["time"] if log else "—",
    }


# ── report ───────────────────────────────────────────────────────────────────

def bar(done: int, total: int, width: int = 10) -> str:
    filled = round(width * done / total) if total else 0
    return "▰" * filled + "▱" * (width - filled)


def pct(done: int, total: int) -> int:
    return round(100 * done / total) if total else 0


def last_note(note: str) -> str:
    return note.strip().split("\n")[-1] if note else ""


def one_line(text: str) -> str:
    return " ".join(text.split())


def short_list(items: list[str], keep: int = 3) -> str:
    if len(items) <= keep + 1:
        return "、".join(items)
    return "、".join(items[:keep]) + f" 等 {len(items)} 項"


def phase_label(name: str) -> str:
    dates = PHASES.get(name, ("", ""))[0]
    return f"{name}（{dates}）" if dates else name


def render_status(p: dict, next_n: int, done_n: int) -> str:
    out = [f"# {PROJECT} 生產就緒進度".replace("#  ", "# "), f"最後更新：{p['last_update']}"]
    if p["current_phase"]:
        goal = PHASES[p["current_phase"]][1]
        out.append(f"目前階段：**{phase_label(p['current_phase'])}**" + (f" — {goal}" if goal else ""))
    out.append("")
    d, t = p["total"]
    out.append(f"**總進度 {d}/{t}（{pct(d, t)}%）** {bar(d, t)}")
    for key, (label, _) in PRIORITIES.items():
        pd_, pt = p["per_priority"][key]
        out.append(f"- {key} {label}：{pd_}/{pt} {bar(pd_, pt)}")
    if PHASES:
        out += ["", "各階段："]
        for ph in PHASES:
            pd_, pt = p["per_phase"][ph]
            mark = " ◀ 目前" if ph == p["current_phase"] else ""
            out.append(f"- {phase_label(ph)}：{pd_}/{pt} {bar(pd_, pt)}{mark}")

    out += ["", "## ✅ 已完成與測試效果"]
    if not p["done"]:
        out.append("（還沒有）")
    for task in p["done"][:done_n]:
        meta = " · ".join(x for x in (task["done_date"], task["commit"] and f"commit {task['commit']}") if x)
        out.append(f"- **{task['id']}** {task['title']}（{meta}）")
        out.append(f"  測試{task['result']}：{one_line(task['evidence']) or '（沒有記錄證據）'}")
    if len(p["done"]) > done_n:
        out.append(f"- …另有 {len(p['done']) - done_n} 項，見「任務清單」")

    out += ["", "## ⏭ 接下來"]
    items = p["in_progress"] + p["ready"][:next_n]
    if not items:
        out.append("（沒有可以開始的任務，看看下面的依賴和待確認問題）")
    for i, task in enumerate(items, start=1):
        tags = [task["priority"], task["phase"], "進行中" if task["status"] == "進行中" else ""]
        out.append(f"{i}. **{task['id']}** {task['title']}〔{' · '.join(x for x in tags if x)}〕")
        if task["status"] == "進行中" and task["result"] != "未測試":
            out.append(f"   上次測試：{task['result']} — {one_line(task['evidence'])}")
        out.append(f"   驗收：{one_line(task['accept'])}")
    # Priority order within a phase would hide other tracks (hardware, content)
    # behind a long run of security items, though they can run side by side.
    parallel = [t["id"] for t in p["ready"][next_n:] if t["phase"] == p["current_phase"]]
    if parallel:
        out.append(f"\n同一階段還可以並行開始：{'、'.join(parallel)}")

    if p["blocked"]:
        out += ["", "## ⛔ 阻塞"]
        for task in p["blocked"]:
            out.append(f"- **{task['id']}** {task['title']} — {last_note(task['note']) or '（未寫原因）'}")

    if p["waiting"]:
        waits = "、".join(f"{t['id']}（等 {short_list(u)}）" for t, u in p["waiting"])
        out += ["", f"⏳ 等待依賴：{waits}"]

    if p["open_questions"]:
        out += ["", f"## ❓ 待確認問題（{len(p['open_questions'])} 個未回覆）"]
        for q in p["open_questions"]:
            out.append(f"- **{q['id']}** {q['question']} → 影響 {q['affects']}")
    return "\n".join(out)


# ── dashboard sheet ──────────────────────────────────────────────────────────

def build_dashboard(wb) -> dict:
    """Rebuilt from scratch on every save; nothing on it is meant to be edited.

    Returns the value each formula cell evaluates to, keyed by cell reference,
    so save() can store them as the formulas' cached results.
    """
    if DASH in wb.sheetnames:
        wb.remove(wb[DASH])
    ws = wb.create_sheet(DASH, 0)
    tasks, questions, log = read_all(wb)
    p = plan(tasks, questions, log)
    cached: dict = {}

    widths = [12, 42, 9, 11, 9, 9, 9, 9, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False

    def put(row, col, value, font=BASE_FONT, align=WRAP_TOP, fill=None, border=True):
        c = ws.cell(row=row, column=col, value=value)
        c.font, c.alignment = font, align
        if fill:
            c.fill = fill
        if border:
            c.border = BORDER
        return c

    def merged(row, first, last, value, font=BASE_FONT, border=True):
        put(row, first, value, font, border=border)
        if last > first:
            for col in range(first + 1, last + 1):
                if border:
                    ws.cell(row=row, column=col).border = BORDER
            ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)

    def header(row, labels):
        for col, label in labels:
            put(row, col, label, HEADER_FONT,
                Alignment(horizontal="center", vertical="center", wrap_text=True), HEADER_FILL)

    def fit(row, spans):
        ws.row_dimensions[row].height = row_height(
            (ws.cell(row=row, column=a).value, sum(widths[a - 1:b])) for a, b in spans
        )

    sta_col = col_letter(TASK_COLS, "status")
    sta = f"'{TASKS}'!${sta_col}:${sta_col}"

    def count_table(row, title, groups, field) -> int:
        """Live COUNTIFS per group and status, so hand edits in Excel show up here."""
        L = col_letter(TASK_COLS, field)
        rng = f"'{TASKS}'!${L}:${L}"
        header(row, list(enumerate([title, "說明", "總數", *STATUSES, "完成率"], start=1)))
        hdr, first = row, row + 1
        totals = [0] * 6

        def ratio(c):
            return c[3] / (c[0] - c[5]) if c[0] - c[5] else 0

        for r, (key, desc) in enumerate(groups, start=first):
            put(r, 1, key, BOLD_FONT, CENTER)
            put(r, 2, desc)
            put(r, 3, f"=COUNTIF({rng},$A{r})", align=CENTER)
            for col in range(4, 9):
                put(r, col, f"=COUNTIFS({rng},$A{r},{sta},{get_column_letter(col)}${hdr})", align=CENTER)
            counts = [sum(t[field] == key for t in tasks)]
            counts += [sum(t[field] == key and t["status"] == s for t in tasks) for s in STATUSES]
            for col, n in enumerate(counts, start=3):
                cached[f"{get_column_letter(col)}{r}"] = n
                totals[col - 3] += n
            put(r, 9, f"=IFERROR(F{r}/(C{r}-H{r}),0)", align=CENTER).number_format = "0%"
            cached[f"I{r}"] = ratio(counts)
            fit(r, [(2, 2)])
        last = first + len(groups)
        put(last, 1, "合計", BOLD_FONT, CENTER)
        put(last, 2, None)
        for col in range(3, 9):
            L2 = get_column_letter(col)
            put(last, col, f"=SUM({L2}{first}:{L2}{last - 1})", BOLD_FONT, CENTER)
            cached[f"{L2}{last}"] = totals[col - 3]
        put(last, 9, f"=IFERROR(F{last}/(C{last}-H{last}),0)", BOLD_FONT, CENTER).number_format = "0%"
        cached[f"I{last}"] = ratio(totals)
        ws.conditional_formatting.add(f"I{first}:I{last}", DataBarRule(
            start_type="num", start_value=0, end_type="num", end_value=1, color="63BE7B"))
        return last

    d, t = p["total"]
    merged(1, 1, 9, f"{PROJECT} 生產就緒追蹤".strip(), Font(name=FONT, size=16, bold=True, color="1F4E79"), border=False)
    ws.row_dimensions[1].height = 26
    merged(2, 1, 9, f"最後更新：{p['last_update']}　·　本頁由 tracker.py 自動生成，請勿手動編輯", MUTED_FONT, border=False)
    summary = f"總進度：{d}/{t} 已完成（{pct(d, t)}%）"
    if p["current_phase"]:
        summary += f"　·　目前階段：{phase_label(p['current_phase'])}"
    merged(3, 1, 9, summary, Font(name=FONT, size=11, bold=True), border=False)

    row = count_table(5, "優先級", [(k, f"{label}：{when}") for k, (label, when) in PRIORITIES.items()],
                      "priority")
    if PHASES:
        groups = [(name, "：".join(x for x in (dates, goal) if x)) for name, (dates, goal) in PHASES.items()]
        row = count_table(row + 2, "階段", groups, "phase")

    row += 2
    merged(row, 1, 9, "⏭ 接下來做什麼", SECTION_FONT, border=False)
    row += 1
    header(row, [(1, "編號"), (2, "任務"), (3, "階段"), (4, "狀態"), (5, "驗收標準 / 上次測試")])
    ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=9)
    items = p["in_progress"] + p["ready"][:6]
    for task in items or [None]:
        row += 1
        if task is None:
            merged(row, 1, 9, "（沒有可以開始的任務）")
            continue
        put(row, 1, task["id"], align=CENTER_TOP)
        put(row, 2, task["title"])
        put(row, 3, task["phase"] or "—", align=CENTER_TOP)
        put(row, 4, task["status"], align=CENTER_TOP)
        detail = one_line(task["accept"])
        if task["status"] == "進行中" and task["result"] != "未測試":
            detail = f"上次測試：{task['result']} — {one_line(task['evidence'])}\n{detail}"
        merged(row, 5, 9, detail)
        fit(row, [(2, 2), (3, 3), (5, 9)])

    row += 2
    merged(row, 1, 9, "✅ 最近完成與測試效果", SECTION_FONT, border=False)
    row += 1
    header(row, [(1, "編號"), (2, "任務"), (3, "測試結果"), (4, "完成日期"), (5, "測試證據")])
    ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=9)
    for task in p["done"][:8] or [None]:
        row += 1
        if task is None:
            merged(row, 1, 9, "（還沒有）")
            continue
        put(row, 1, task["id"], align=CENTER_TOP)
        put(row, 2, task["title"])
        put(row, 3, task["result"], align=CENTER_TOP)
        put(row, 4, task["done_date"], align=CENTER_TOP)
        merged(row, 5, 9, one_line(task["evidence"]))
        fit(row, [(2, 2), (5, 9)])

    if p["waiting"] or p["blocked"]:
        row += 2
        merged(row, 1, 9, "⏳ 等待依賴 / ⛔ 阻塞", SECTION_FONT, border=False)
        row += 1
        header(row, [(1, "編號"), (2, "任務"), (3, "在等什麼")])
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=9)
        for task in p["blocked"]:
            row += 1
            put(row, 1, task["id"], align=CENTER_TOP)
            put(row, 2, task["title"])
            merged(row, 3, 9, f"阻塞：{last_note(task['note']) or '（未寫原因）'}")
            fit(row, [(2, 2), (3, 9)])
        for task, unmet in p["waiting"]:
            row += 1
            put(row, 1, task["id"], align=CENTER_TOP)
            put(row, 2, task["title"])
            merged(row, 3, 9, "等 " + "、".join(unmet))
            fit(row, [(2, 2), (3, 9)])

    if p["open_questions"]:
        row += 2
        merged(row, 1, 9, "❓ 待確認問題", SECTION_FONT, border=False)
        row += 1
        header(row, [(1, "編號"), (2, "問題"), (3, "影響的任務 / 為什麼重要")])
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=9)
        for q in p["open_questions"]:
            row += 1
            put(row, 1, q["id"], align=CENTER_TOP)
            put(row, 2, q["question"])
            merged(row, 3, 9, f"{q['affects']}：{q['why']}")
            fit(row, [(2, 2), (3, 9)])

    row += 2
    merged(row, 1, 9, "說明", SECTION_FONT, border=False)
    notes = [
        "狀態：未開始 → 進行中 → 已完成。只有測試真的跑過而且通過才算已完成；阻塞要在備註寫原因；不再需要的任務標不適用，不要刪。",
        "測試結果：未測試 / 通過 / 失敗 / 部分通過。測試證據寫跑了什麼、看到什麼，完整歷史在「更新記錄」。",
        "「接下來」按階段排序，同一階段內再按優先級；依賴還沒滿足的任務不會出現在「接下來」。",
        "可以手動改：「任務清單」的狀態、測試結果、測試證據、Commit、備註，以及「待確認問題」的回答。其餘欄位和本頁由腳本維護。",
        "建議用法：在 Claude Code 對 Claude 說「P0-02 做完了」或輸入 /prod-readiness。Claude 會先跑測試再更新，並寫進「更新記錄」。",
    ]
    for note in notes:
        row += 1
        merged(row, 1, 9, "・" + note, border=False)
        fit(row, [(1, 9)])
    return cached


def store_cached_values(xlsx: str, values: dict) -> None:
    """Give the dashboard's formulas cached results.

    openpyxl writes formulas with an empty <v>, and anything that does not
    recalculate (Quick Look, file previews, pandas) then shows every count as 0.
    Excel still recalculates on open because of fullCalcOnLoad.
    """
    sheet = "xl/worksheets/sheet1.xml"  # openpyxl numbers sheets in tab order
    with zipfile.ZipFile(xlsx) as zin:
        entries = [(info, zin.read(info.filename)) for info in zin.infolist()]
    xml = dict((i.filename, d) for i, d in entries)[sheet].decode("utf-8")
    if "<f>COUNTIF(" not in xml:
        return  # not the dashboard; leave the file as openpyxl wrote it

    def fill(m: re.Match) -> str:
        ref = m.group(2)
        if ref not in values:
            return m.group(0)
        return f"{m.group(1)}<f>{m.group(3)}</f><v>{values[ref]}</v></c>"

    xml = re.sub(r'(<c r="([A-Z]+\d+)"[^>]*>)<f>(.*?)</f><v\s*/?>(?:</v>)?</c>', fill, xml)
    with zipfile.ZipFile(xlsx, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in entries:
            zout.writestr(info, xml.encode("utf-8") if info.filename == sheet else data)


# ── persistence ──────────────────────────────────────────────────────────────

def excel_has_it_open(path: Path) -> bool:
    # Excel's owner file is "~$" + name, except on Windows for long names,
    # where it replaces the first two characters instead.
    return any((path.parent / n).exists() for n in ("~$" + path.name, "~$" + path.name[2:]))


def open_book(path: Path, for_write: bool, allow_open: bool = False):
    if not path.exists():
        die(f"找不到 {path}。先執行：tracker.py init")
    if for_write and not allow_open and excel_has_it_open(path):
        die(f"{path.name} 正在 Excel 裡打開。請先關閉（否則你在 Excel 保存時會覆蓋這次更新），"
            "或加 --ignore-lock。")
    return load_workbook(path)


def backup(path: Path) -> Path:
    target = path.with_name(f".{path.stem}.{dt.datetime.now():%Y%m%d-%H%M%S}.bak.xlsx")
    shutil.copy2(path, target)
    return target


def save(wb, path: Path) -> None:
    cached = build_dashboard(wb)
    for name, cols in ((TASKS, TASK_COLS), (QUESTIONS, QUESTION_COLS), (LOG, LOG_COLS)):
        ws = wb[name]
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(ws.max_row, 2)}"
    # No cached values exist for openpyxl-written formulas; this makes Excel
    # and Numbers compute the dashboard counts on open.
    wb.calculation.fullCalcOnLoad = True
    # Exactly one selected tab, or Excel opens with the sheets grouped.
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = False
    wb.active = 0
    wb.worksheets[0].sheet_view.tabSelected = True

    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=path.parent, prefix=".", suffix=".xlsx.tmp")
    os.close(fd)
    try:
        wb.save(tmp)
        store_cached_values(tmp, cached)
        # mkstemp creates 0600; keep the file readable the way it was.
        os.chmod(tmp, path.stat().st_mode & 0o777 if path.exists() else 0o644)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)


def append_log(wb, **entry) -> None:
    ws = wb[LOG]
    cmap = header_map(ws, LOG_COLS)
    rows, _ = read_rows(ws, LOG_COLS)
    row = (rows[-1]["_row"] + 1) if rows else 2
    write_cells(ws, row, cmap, LOG_COLS, {"time": now_str(), **entry})


def find(rows, rid: str, what: str):
    rid = rid.upper()
    for r in rows:
        if r["id"] == rid:
            return r
    die(f"找不到{what} {rid}")


def norm(value: str | None, allowed, aliases, what: str) -> str | None:
    if value is None:
        return None
    v = aliases.get(value.strip().lower(), value.strip())
    if v not in allowed:
        die(f"{what}必須是：{'、'.join(allowed)}（收到「{value}」）")
    return v


def norm_phase(value: str | None) -> str | None:
    """Accept a phase's full name or just its leading number ("2")."""
    if value is None:
        return None
    if not PHASES:
        die("seed.json 沒有定義任何階段。")
    v = value.strip()
    if v in PHASES:
        return v
    matches = [name for name in PHASES if name.split()[0] == v]
    if len(matches) == 1:
        return matches[0]
    die(f"階段必須是：{'、'.join(PHASES)}（收到「{value}」）")


# ── commands ─────────────────────────────────────────────────────────────────

def cmd_init(args) -> None:
    path = Path(args.file)
    if path.exists() and not args.force:
        die(f"{path} 已存在。要重建（會清掉所有狀態）請加 --force；舊檔會備份。"
            "只想帶入 seed.json 的新任務或新版面，用 tracker.py sync。")
    if path.exists():
        print(f"已備份舊檔到 {backup(path)}")
    wb = build_workbook(SEED["tasks"], SEED["questions"], [])
    append_log(wb, id="—", title="建立追蹤表",
               note=f"依據 {SEED.get('review_date', today())} 的評估建立："
                    f"{len(SEED['tasks'])} 項任務、{len(SEED['questions'])} 個待確認問題")
    save(wb, path)
    print(f"已建立 {path}")


def cmd_sync(args) -> None:
    """Rebuild in the current layout, keeping every row, and add what seed.json has new.

    Existing rows are never overwritten; the only field filled in is an empty
    phase, from the same task in seed.json.
    """
    path = Path(args.file)
    wb = open_book(path, for_write=True, allow_open=args.ignore_lock)
    tasks, questions, log = read_all(wb, allow_missing=True)

    seed_tasks = {t["id"].upper(): t for t in SEED["tasks"]}
    phased = []
    for t in tasks:
        seed_phase = seed_tasks.get(t["id"], {}).get("phase", "")
        if not t["phase"] and seed_phase:
            t["phase"] = seed_phase
            phased.append(t["id"])
    have = {t["id"] for t in tasks}
    added = [dict(t, id=tid) for tid, t in seed_tasks.items() if tid not in have]
    tasks += added
    q_have = {q["id"] for q in questions}
    q_added = [q for q in SEED["questions"] if q["id"].upper() not in q_have]
    questions += q_added

    print(f"已備份舊檔到 {backup(path)}")
    new_wb = build_workbook(tasks, questions, log)
    for t in added:
        append_log(new_wb, id=t["id"], title=t["title"], new_status="未開始",
                   note="新增任務（從 seed.json 同步）")
    for q in q_added:
        append_log(new_wb, id=q["id"], title=q["question"], new_status="待回覆",
                   note="新增待確認問題（從 seed.json 同步）")
    append_log(new_wb, id="—", title="同步表格",
               note=f"版面更新到最新版；新增 {len(added)} 項任務、{len(q_added)} 個問題；"
                    f"為 {len(phased)} 項任務填入階段")
    save(new_wb, path)
    print(f"已同步 {path.name}：新增任務 {len(added)} 項"
          + (f"（{short_list([t['id'] for t in added], 5)}）" if added else "")
          + f"，新增問題 {len(q_added)} 個，填入階段 {len(phased)} 項")


def cmd_status(args) -> None:
    wb = open_book(Path(args.file), for_write=False)
    print(render_status(plan(*read_all(wb)), args.next, args.done))


def cmd_show(args) -> None:
    wb = open_book(Path(args.file), for_write=False)
    tasks, questions, log = read_all(wb)
    rid = args.id.upper()
    if rid.startswith("Q"):
        q = find(questions, rid, "問題")
        print(f"**{q['id']} {q['question']}**\n狀態：{q['status']}\n為什麼重要：{q['why']}\n"
              f"影響的任務：{q['affects']}\n回答：{q['answer'] or '—'}（{q['answered_on'] or '—'}）")
        return
    t = find(tasks, rid, "任務")
    unmet = unmet_deps(t, {x["id"]: x for x in tasks}, {q["id"]: q for q in questions})
    deps = t["deps"] or "—"
    if unmet:
        deps += f"（未滿足：{'、'.join(unmet)}）"
    print(f"**{t['id']} {t['title']}**")
    print(f"優先級：{t['priority']} · 階段：{phase_label(t['phase']) if t['phase'] else '—'} · "
          f"類別：{t['category']} · 狀態：{t['status']} · 測試結果：{t['result']} · "
          f"完成日期：{t['done_date'] or '—'}")
    print(f"依賴：{deps}")
    for key, label in (("risk", "問題 / 風險"), ("where", "相關位置"), ("accept", "驗收標準"),
                       ("test", "測試方法"), ("evidence", "測試證據"), ("commit", "Commit"),
                       ("note", "備註")):
        print(f"\n{label}：\n{t[key] or '—'}")
    history = [e for e in log if e["id"].upper() == rid]
    print("\n更新記錄：")
    if not history:
        print("—")
    for e in history:
        change = (e["new_status"] if e["old_status"] == e["new_status"]
                  else f"{e['old_status'] or '—'} → {e['new_status'] or '—'}")
        extra = " · ".join(x for x in (e["result"], one_line(e["evidence"]), e["commit"], one_line(e["note"])) if x)
        print(f"- {e['time']} {change}" + (f" · {extra}" if extra else ""))


def cmd_list(args) -> None:
    wb = open_book(Path(args.file), for_write=False)
    tasks, _, _ = read_all(wb)
    status = norm(args.status, STATUSES, STATUS_ALIASES, "狀態")
    phase = norm_phase(args.phase)
    rows = [t for t in sorted(tasks, key=sort_key)
            if (not args.priority or t["priority"] == args.priority.upper())
            and (not phase or t["phase"] == phase)
            and (not status or t["status"] == status)]
    print("| 編號 | 階段 | 狀態 | 測試結果 | 任務 |\n|---|---|---|---|---|")
    for t in rows:
        print(f"| {t['id']} | {t['phase'] or '—'} | {t['status']} | {t['result']} | {t['title']} |")
    print(f"\n共 {len(rows)} 項")


def cmd_update(args) -> None:
    path = Path(args.file)
    wb = open_book(path, for_write=True, allow_open=args.ignore_lock)
    tasks, questions, _ = read_all(wb)
    t = find(tasks, args.id, "任務")
    status = norm(args.status, STATUSES, STATUS_ALIASES, "狀態")
    result = norm(args.result, RESULTS, RESULT_ALIASES, "測試結果")
    phase = norm_phase(args.phase)

    if status == "已完成":
        if result in ("失敗", "部分通過"):
            die("測試沒有完全通過，不能標記已完成。去掉 --status，保持進行中。")
        if args.force:
            if not args.note:
                die("--force 需要 --note 說明為什麼不經測試就標記完成。")
        else:
            if result != "通過":
                die("標記已完成必須同時給 --result 通過：這次實際跑過、而且通過的測試。")
            if not args.evidence:
                die("標記已完成必須給 --evidence：跑了什麼、看到什麼。")
    if result in ("失敗", "部分通過") and status is None and t["status"] == "未開始":
        status = "進行中"

    changes: dict[str, str] = {}
    if status and status != t["status"]:
        changes["status"] = status
        if status == "已完成":
            changes["done_date"] = today()
        elif t["status"] == "已完成":
            changes["done_date"] = ""
    for key, value in (("result", result), ("evidence", args.evidence), ("commit", args.commit),
                       ("test", args.test_method), ("accept", args.accept),
                       ("phase", phase), ("deps", args.deps)):
        if value is not None:
            changes[key] = value
    if args.note:
        entry = f"[{today()}] {args.note}"
        changes["note"] = f"{t['note']}\n{entry}" if t["note"] else entry
    if not changes:
        die("沒有要更新的內容。")

    # Edits to the plan itself (criteria, phase, dependencies) go in the log
    # too, so nobody can quietly move the goalposts.
    labels = {k: label for k, label, _ in TASK_COLS}
    edited = [f"{labels[k]}→{changes[k]}" if k in ("phase", "deps") else labels[k]
              for k in ("accept", "test", "phase", "deps") if k in changes]
    log_note = "；".join(x for x in (args.note or "", "修改了" + "、".join(edited) if edited else "") if x)

    ws = wb[TASKS]
    write_cells(ws, t["_row"], header_map(ws, TASK_COLS), TASK_COLS, changes)
    append_log(wb, id=t["id"], title=t["title"], old_status=t["status"],
               new_status=changes.get("status", t["status"]), result=result or "",
               evidence=args.evidence or "", commit=args.commit or "", note=log_note)
    save(wb, path)

    moved = (f"{t['status']} → {changes['status']}" if "status" in changes
             else f"狀態維持{t['status']}")
    print(f"已更新 {t['id']} {t['title']}：{moved}" + (f"，測試{result}" if result else "")
          + (f"；{'、'.join(edited)}" if edited else ""))
    if changes.get("status") == "已完成":
        unmet = unmet_deps(t, {x["id"]: x for x in tasks}, {q["id"]: q for q in questions})
        if unmet:
            print(f"注意：它的依賴 {'、'.join(unmet)} 還沒完成。")


def cmd_answer(args) -> None:
    path = Path(args.file)
    wb = open_book(path, for_write=True, allow_open=args.ignore_lock)
    tasks, questions, _ = read_all(wb)
    q = find(questions, args.id, "問題")
    ws = wb[QUESTIONS]
    write_cells(ws, q["_row"], header_map(ws, QUESTION_COLS), QUESTION_COLS,
                {"answer": args.text, "answered_on": today(), "status": "已回覆"})
    append_log(wb, id=q["id"], title=q["question"], old_status=q["status"],
               new_status="已回覆", note=args.text)
    save(wb, path)
    print(f"已記錄 {q['id']} 的回答。")

    q["status"] = "已回覆"
    by_id = {t["id"]: t for t in tasks}
    q_by_id = {x["id"]: x for x in questions}
    freed = [t for t in sorted(tasks, key=sort_key)
             if t["status"] == "未開始" and q["id"] in split_deps(t["deps"])
             and not unmet_deps(t, by_id, q_by_id)]
    if freed:
        print("現在可以開始：" + "、".join(f"{t['id']} {t['title']}" for t in freed))


def cmd_ask(args) -> None:
    path = Path(args.file)
    wb = open_book(path, for_write=True, allow_open=args.ignore_lock)
    _, questions, _ = read_all(wb)
    qid = f"Q{max((id_num(q['id']) for q in questions), default=0) + 1}"
    ws = wb[QUESTIONS]
    row = (max(q["_row"] for q in questions) + 1) if questions else 2
    write_cells(ws, row, header_map(ws, QUESTION_COLS), QUESTION_COLS,
                {"id": qid, "question": args.question, "why": args.why or "",
                 "affects": args.affects or "", "status": "待回覆"})
    append_log(wb, id=qid, title=args.question, new_status="待回覆", note="新增待確認問題")
    save(wb, path)
    print(f"已新增 {qid} {args.question}")


def cmd_add(args) -> None:
    path = Path(args.file)
    wb = open_book(path, for_write=True, allow_open=args.ignore_lock)
    tasks, _, _ = read_all(wb)
    priority = args.priority.upper()
    if priority not in PRIORITIES:
        die(f"優先級必須是：{'、'.join(PRIORITIES)}")
    n = max((id_num(t["id"]) for t in tasks if t["priority"] == priority), default=0) + 1
    new_id = f"{priority}-{n:02d}"
    rec = {"id": new_id, "priority": priority, "phase": norm_phase(args.phase) or "",
           "category": args.category or "", "title": args.title, "status": "未開始",
           "result": "未測試", "risk": args.risk or "", "where": args.where or "",
           "accept": args.accept or "", "test": args.test or "", "deps": args.deps or "",
           "note": f"[{today()}] 新增" + (f"：{args.note}" if args.note else "")}
    ws = wb[TASKS]
    row = (max(t["_row"] for t in tasks) + 1) if tasks else 2
    write_cells(ws, row, header_map(ws, TASK_COLS), TASK_COLS, rec)
    append_log(wb, id=new_id, title=args.title, new_status="未開始", note="新增任務")
    save(wb, path)
    print(f"已新增 {new_id} {args.title}")


def main() -> None:
    parser = argparse.ArgumentParser(description=f"{PROJECT} 生產就緒追蹤".strip())
    parser.add_argument("--file", default=str(DEFAULT_FILE), help="追蹤表路徑")
    sub = parser.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("init", help="從 assets/seed.json 建立全新的追蹤表")
    p.add_argument("--force", action="store_true")
    p.set_defaults(func=cmd_init)

    p = sub.add_parser("sync", help="把 seed.json 新增的任務、問題和最新版面帶進現有追蹤表")
    p.add_argument("--ignore-lock", action="store_true")
    p.set_defaults(func=cmd_sync)

    p = sub.add_parser("status", help="進度報告")
    p.add_argument("--next", type=int, default=3, help="列出幾項接下來可以開始的任務")
    p.add_argument("--done", type=int, default=5, help="列出幾項最近完成的任務")
    p.set_defaults(func=cmd_status)

    p = sub.add_parser("show", help="顯示一項任務或問題的全部內容")
    p.add_argument("id")
    p.set_defaults(func=cmd_show)

    p = sub.add_parser("list", help="任務清單")
    p.add_argument("--priority")
    p.add_argument("--phase")
    p.add_argument("--status")
    p.set_defaults(func=cmd_list)

    for name in ("update", "start"):
        p = sub.add_parser(name, help="更新任務" if name == "update" else "把任務標記為進行中")
        p.add_argument("id")
        if name == "update":
            p.add_argument("--status")
            p.add_argument("--result")
            p.add_argument("--evidence")
            p.add_argument("--commit")
            p.add_argument("--test-method", dest="test_method")
            p.add_argument("--accept")
            p.add_argument("--phase")
            p.add_argument("--deps")
            p.add_argument("--force", action="store_true", help="不經測試標記完成（需 --note）")
        p.add_argument("--note")
        p.add_argument("--ignore-lock", action="store_true")
        p.set_defaults(func=cmd_update)
        if name == "start":
            p.set_defaults(status="進行中", result=None, evidence=None, commit=None,
                           test_method=None, accept=None, phase=None, deps=None, force=False)

    p = sub.add_parser("answer", help="記錄待確認問題的回答")
    p.add_argument("id")
    p.add_argument("--text", required=True)
    p.add_argument("--ignore-lock", action="store_true")
    p.set_defaults(func=cmd_answer)

    p = sub.add_parser("ask", help="新增待確認問題")
    p.add_argument("--question", required=True)
    p.add_argument("--why")
    p.add_argument("--affects")
    p.add_argument("--ignore-lock", action="store_true")
    p.set_defaults(func=cmd_ask)

    p = sub.add_parser("add", help="新增任務")
    p.add_argument("--priority", required=True)
    p.add_argument("--title", required=True)
    for opt in ("phase", "category", "risk", "where", "accept", "test", "deps", "note"):
        p.add_argument(f"--{opt}")
    p.add_argument("--ignore-lock", action="store_true")
    p.set_defaults(func=cmd_add)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
